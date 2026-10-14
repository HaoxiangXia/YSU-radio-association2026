import io
import json
import zipfile
from datetime import datetime, timedelta, timezone

import jwt
import pytest
from openpyxl import Workbook

from conftest import TEST_PASSWORD
from models.admission_list import AdmissionDataError
from routes import recruitment_officers
from utils.admission_workbook import create_admission_template, parse_admission_workbook
from utils.security import admission_query_limiter, login_limiter


def login(client):
    response = client.post(
        "/api/recruitment-officers/login",
        json={"username": "officer", "password": TEST_PASSWORD},
    )
    assert response.status_code == 200
    # 会话经 Set-Cookie 写入 TestClient 的 cookie jar，后续请求自动携带


def workbook_bytes(*, formula: bool = False, extra_header: bool = False) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "录取名单"
    headers = ["姓名", "学号"]
    if extra_header:
        headers.append("备注")
    worksheet.append(headers)
    worksheet.append((
        "=\"新同学\"" if formula else "新同学",
        "202600000008",
    ))
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_login_success_verify_and_wrong_password(default_client):
    client, _ = default_client

    wrong = client.post(
        "/api/recruitment-officers/login",
        json={"username": "officer", "password": "wrong"},
    )
    success = client.post(
        "/api/recruitment-officers/login",
        json={"username": "officer", "password": TEST_PASSWORD},
    )
    verified = client.get("/api/recruitment-officers/verify")

    assert wrong.status_code == 401
    assert success.status_code == 200
    assert "token" not in success.json()
    set_cookie = success.headers["set-cookie"].lower()
    assert "httponly" in set_cookie
    assert "secure" in set_cookie
    assert "samesite=strict" in set_cookie
    assert verified.status_code == 200
    assert verified.json()["officer"]["username"] == "officer"


def test_login_over_http_omits_secure_flag(client_factory):
    # 内网/本地 HTTP 访问：Secure Cookie 会被浏览器拒存，登录后被弹回登录页（回归）
    with client_factory(base_url="http://testserver") as (client, _):
        success = client.post(
            "/api/recruitment-officers/login",
            json={"username": "officer", "password": TEST_PASSWORD},
        )
        verified = client.get("/api/recruitment-officers/verify")

        assert success.status_code == 200
        set_cookie = success.headers["set-cookie"].lower()
        assert "httponly" in set_cookie
        assert "secure" not in set_cookie
        assert verified.status_code == 200


def test_successful_login_clears_failure_budget(default_client):
    client, _ = default_client
    payload = {"username": "officer", "password": "wrong"}
    for _ in range(4):
        assert client.post("/api/recruitment-officers/login", json=payload).status_code == 401

    assert client.post(
        "/api/recruitment-officers/login",
        json={"username": "officer", "password": TEST_PASSWORD},
    ).status_code == 200

    for _ in range(5):
        assert client.post("/api/recruitment-officers/login", json=payload).status_code == 401
    assert client.post("/api/recruitment-officers/login", json=payload).status_code == 429
    login_limiter.clear("testclient")


def test_expired_and_invalid_tokens_are_rejected(default_client):
    client, _ = default_client
    expired = jwt.encode(
        {
            "recruitmentOfficerId": "officer",
            "username": "officer",
            "name": "officer",
            "exp": datetime.now(timezone.utc) - timedelta(seconds=1),
        },
        recruitment_officers.JWT_SECRET,
        algorithm="HS256",
    )

    client.cookies.set(recruitment_officers.SESSION_COOKIE_NAME, expired)
    expired_response = client.get("/api/recruitment-officers/verify")
    client.cookies.set(recruitment_officers.SESSION_COOKIE_NAME, "not-a-token")
    invalid_response = client.get("/api/recruitment-officers/verify")

    assert expired_response.status_code == 401
    assert invalid_response.status_code == 403


def test_admission_query_success_and_privacy(default_client):
    client, _ = default_client
    response = client.post(
        "/api/admissions/query",
        json={"studentId": "202600000001"},
    )

    assert response.status_code == 200
    assert response.json() == {
        "name": "测试同学",
        "department": "",
        "status": "已录取",
    }
    assert "studentId" not in response.json()
    assert "phone" not in response.json()


def test_unknown_student_has_no_result(default_client):
    client, _ = default_client
    response = client.post("/api/admissions/query", json={"studentId": "202600000009"})
    assert response.status_code == 404
    assert "未找到匹配的录取结果" in response.json()["detail"]


@pytest.mark.parametrize("payload", [
    {}, {"studentId": "123"}, {"studentId": "abcdefghijkl"},
    {"studentId": "202600000001", "phone": "13800000001"},
])
def test_query_rejects_invalid_student_id_and_legacy_phone(default_client, payload):
    client, _ = default_client
    assert client.post("/api/admissions/query", json=payload).status_code == 422


def test_query_trims_student_id(default_client):
    client, _ = default_client
    assert client.post("/api/admissions/query", json={"studentId": " 202600000001 "}).status_code == 200


def test_admission_query_closed(client_factory, config_copy):
    config = config_copy(admission_enabled=False)
    with client_factory(config=config) as (client, _):
        response = client.post(
            "/api/admissions/query",
            json={"studentId": "202600000001"},
        )

    assert response.status_code == 403
    assert response.json()["detail"] == "录取查询尚未开放，请留意协会后续通知。"


def test_admission_rate_limit(default_client):
    client, _ = default_client
    payload = {"studentId": "202600000009"}

    for _ in range(20):
        assert client.post("/api/admissions/query", json=payload).status_code == 404
    assert client.post("/api/admissions/query", json=payload).status_code == 429
    admission_query_limiter.clear("testclient")


def test_private_admissions_file_is_not_public(default_client):
    client, _ = default_client
    for path in [
        "/data/admission-results.json",
        "/admission-results.json",
        "/config/recruitment.local.json",
    ]:
        assert client.get(path).status_code == 404


def test_invalid_admissions_file_stops_startup_without_echoing_private_value(
    client_factory,
):
    private_value = "13912345678"
    invalid_records = [
        {
            "studentId": "202600000001",
            "name": "测试同学",
            "phone": private_value,
            "department": "",
            "status": "待定",
        }
    ]

    try:
        with client_factory(admissions=invalid_records):
            pass
    except AdmissionDataError as error:
        assert private_value not in str(error)
    else:
        raise AssertionError("无效录取名单没有阻止应用启动")


def test_admissions_management_requires_authentication(default_client):
    client, _ = default_client

    assert client.get("/api/admissions/manage/status").status_code == 401
    assert client.get("/api/admissions/manage/template.xlsx").status_code == 401
    assert client.post(
        "/api/admissions/manage/preview",
        content=workbook_bytes(),
    ).status_code == 401


def test_officer_can_preview_publish_and_then_open_query(client_factory, config_copy):
    config = config_copy(admission_enabled=False)
    with client_factory(config=config, admissions=[]) as (client, state):
        login(client)
        template = client.get("/api/admissions/manage/template.xlsx")
        preview = client.post(
            "/api/admissions/manage/preview",
            headers={
                "Content-Type": (
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                ),
            },
            content=workbook_bytes(),
        )
        publish = client.post(
            "/api/admissions/manage/publish",
            json={"previewId": preview.json()["previewId"]},
        )
        config["admissionQuery"]["enabled"] = True
        opened = client.put(
            "/api/recruitment/manage/config",
            json=config,
        )
        query = client.post(
            "/api/admissions/query",
            json={"studentId": "202600000008"},
        )
        persisted = json.loads(state["admissions_path"].read_text(encoding="utf-8"))

    assert template.status_code == 200
    assert template.content.startswith(b"PK")
    assert preview.status_code == 200
    assert preview.json()["valid"] is True
    assert preview.json()["preview"][0] == {"name": "新**", "studentId": "********0008", "status": "已录取"}
    assert publish.status_code == 200
    assert opened.status_code == 200
    assert query.status_code == 200
    assert query.json()["name"] == "新同学"
    assert persisted[0] == {
        "name": "新同学", "studentId": "202600000008", "department": "", "status": "已录取",
    }
    assert query.json() == {"name": "新同学", "department": "", "status": "已录取"}
    assert parse_admission_workbook(template.content)["records"][0].student_id == "202600000001"


def test_formula_is_rejected_and_publish_requires_closed_query(default_client):
    client, _ = default_client
    login(client)
    invalid_preview = client.post(
        "/api/admissions/manage/preview",
        content=workbook_bytes(formula=True),
    )
    valid_preview = client.post(
        "/api/admissions/manage/preview",
        content=workbook_bytes(),
    )
    publish = client.post(
        "/api/admissions/manage/publish",
        json={"previewId": valid_preview.json()["previewId"]},
    )

    assert invalid_preview.status_code == 200
    assert invalid_preview.json()["valid"] is False
    assert invalid_preview.json()["errors"][0]["message"] == "不得使用公式"
    assert publish.status_code == 409
    assert "关闭录取查询" in publish.json()["detail"]


def test_extra_excel_columns_and_expanded_workbook_are_rejected(default_client):
    client, _ = default_client
    login(client)
    extra_column = client.post(
        "/api/admissions/manage/preview",
        content=workbook_bytes(extra_header=True),
    )

    expanded = io.BytesIO()
    with zipfile.ZipFile(expanded, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", b"0" * (21 * 1024 * 1024))

    with pytest.raises(AdmissionDataError, match="展开后过大"):
        parse_admission_workbook(expanded.getvalue())

    assert extra_column.status_code == 422
    assert "只能包含这些表头" in extra_column.json()["detail"]


def test_logout_revokes_session_on_server(default_client):
    client, _ = default_client
    login(client)
    token = client.cookies.get(recruitment_officers.SESSION_COOKIE_NAME)
    assert client.get("/api/recruitment-officers/verify").status_code == 200

    assert client.post("/api/recruitment-officers/logout").status_code == 200

    # 注销后原会话被服务端吊销：把原 cookie 放回去也必须 401
    client.cookies.set(recruitment_officers.SESSION_COOKIE_NAME, token)
    assert client.get("/api/recruitment-officers/verify").status_code == 401


def test_api_docs_are_disabled_in_production(default_client):
    client, _ = default_client
    assert client.get("/docs").status_code == 404
    assert client.get("/redoc").status_code == 404
    assert client.get("/openapi.json").status_code == 404


def test_cross_site_state_changing_requests_are_rejected(default_client):
    client, _ = default_client
    payload = {"studentId": "202600000001"}

    cross_site_origin = client.post(
        "/api/admissions/query",
        json=payload,
        headers={"Origin": "https://untrusted.example"},
    )
    cross_site_fetch = client.post(
        "/api/admissions/query",
        json=payload,
        headers={"Sec-Fetch-Site": "cross-site"},
    )
    same_site = client.post(
        "/api/admissions/query",
        json=payload,
        headers={"Origin": "http://testserver", "Sec-Fetch-Site": "same-origin"},
    )

    assert cross_site_origin.status_code == 403
    assert cross_site_fetch.status_code == 403
    assert same_site.status_code == 200


@pytest.mark.parametrize("headers,rows,error", [
    (["姓名", "学号", "申请手机号", "录取部门"], [["测试", "202600000001", "13800000001", "机械"]], "表头"),
    (["学号", "姓名"], [["202600000001", "测试"]], "表头"),
    (["姓名", "学号"], [], "没有数据行"),
])
def test_workbook_rejects_legacy_headers_wrong_order_and_empty_list(headers, rows, error):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "录取名单"
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    with pytest.raises(AdmissionDataError, match=error):
        parse_admission_workbook(output.getvalue())


@pytest.mark.parametrize("rows,message", [
    ([["同名", "202600000001"], ["同名", "202600000001"]], "学号重复"),
    ([["", "202600000001"]], "字段缺失或格式不正确"),
    ([["测试", "123"]], "字段缺失或格式不正确"),
    ([["长" * 31, "202600000001"]], "字段缺失或格式不正确"),
    ([["测试", "202600000001", "额外内容"]], "只能填写姓名、学号两列"),
])
def test_workbook_reports_invalid_rows(rows, message):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "录取名单"
    worksheet.append(["姓名", "学号"])
    for row in rows:
        worksheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    result = parse_admission_workbook(output.getvalue())
    assert result["errors"][0]["message"] == message


def test_workbook_allows_same_names_blank_rows_and_numeric_student_ids():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "录取名单"
    worksheet.append(["姓名", "学号"])
    worksheet.append([" 同名 ", 202600000001])
    worksheet.append([None, None])
    worksheet.append(["同名", "202600000002"])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    result = parse_admission_workbook(output.getvalue())
    assert not result["errors"]
    assert [r.student_id for r in result["records"]] == ["202600000001", "202600000002"]
    assert all(r.name == "同名" and r.status == "已录取" and r.department == "" for r in result["records"])


def test_workbook_rejects_excessive_rows_and_bytes():
    from utils.admission_workbook import MAX_WORKBOOK_BYTES
    with pytest.raises(AdmissionDataError, match="2 MiB"):
        parse_admission_workbook(b"0" * (MAX_WORKBOOK_BYTES + 1))
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "录取名单"
    worksheet.append(["姓名", "学号"])
    for index in range(2001):
        worksheet.append(["测试", str(202600000000 + index)])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    with pytest.raises(AdmissionDataError, match="最多允许 2000 行"):
        parse_admission_workbook(output.getvalue())


def test_template_contains_only_name_and_student_id():
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(create_admission_template()))
    worksheet = workbook["录取名单"]
    assert worksheet.max_column == 2
    assert tuple(cell.value for cell in worksheet[1]) == ("姓名", "学号")
    assert worksheet.auto_filter.ref == "A1:B2"
    workbook.close()
